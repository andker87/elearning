import csv, io, os
from datetime import datetime, date
from fastapi import HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4, A3, landscape
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
from . import cifratura

FORMATI = {"csv", "xlsx", "pdf"}
PAGINE = {"A4", "A3"}
MEDIA_TYPE = {
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}

def _cella(v, neutra=True):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Sì" if v else ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    return neutralizza(v) if neutra else v  # nel PDF non c'e' nessuna formula da evitare: l'apice si vedrebbe

def neutralizza(v):
    """Un testo che inizia con = + - @ verrebbe eseguito come formula da Excel/LibreOffice: lo si rende testo."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v

def _a_csv(colonne, righe) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(colonne)
    for r in righe:
        w.writerow([_cella(v) for v in r])
    return buf.getvalue().encode("utf-8-sig")

def _a_excel(titolo, colonne, righe) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = (titolo or "Foglio1")[:31]
    ws.append(colonne)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in righe:
        ws.append([_cella(v) for v in r])
    for i, col in enumerate(colonne, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(12, min(40, len(str(col)) + 4))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def _a_pdf(titolo, colonne, righe, pagina="A4", logo_path=None) -> bytes:
    pagesize = landscape(A3 if pagina == "A3" else A4)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=pagesize, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
                             leftMargin=1.5 * cm, rightMargin=1.5 * cm)
    stili = getSampleStyleSheet()
    elementi = []
    if logo_path and os.path.exists(logo_path):
        try:
            elementi.append(Image(io.BytesIO(cifratura.leggi_file(logo_path)), width=3 * cm, height=1.3 * cm, kind="proportional"))
            elementi.append(Spacer(1, 0.3 * cm))
        except Exception:
            pass
    elementi.append(Paragraph(titolo, stili["Title"]))
    elementi.append(Spacer(1, 0.4 * cm))
    dati = [colonne] + [[_cella(v, neutra=False) for v in r] for r in righe]
    tabella = Table(dati, repeatRows=1)
    tabella.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2d3d")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f5f5")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    elementi.append(tabella)
    doc.build(elementi)
    return buf.getvalue()

def esporta(formato: str, nome_file: str, titolo: str, colonne, righe, pagina: str = "A4", logo_path: str = None) -> StreamingResponse:
    """Esporta una tabella (colonne + righe come liste di valori, non dict)
    in CSV, Excel o PDF (A4/A3). Punto unico per tutte le esportazioni della
    piattaforma, così formati e layout restano coerenti ovunque."""
    if formato not in FORMATI:
        raise HTTPException(400, f"Formato non valido: {formato} (usa csv, xlsx o pdf)")
    if pagina not in PAGINE:
        raise HTTPException(400, f"Formato pagina non valido: {pagina} (usa A4 o A3)")
    if formato == "csv":
        contenuto = _a_csv(colonne, righe)
    elif formato == "xlsx":
        contenuto = _a_excel(titolo, colonne, righe)
    else:
        contenuto = _a_pdf(titolo, colonne, righe, pagina, logo_path)
    return StreamingResponse(
        iter([contenuto]), media_type=MEDIA_TYPE[formato],
        headers={"Content-Disposition": f'attachment; filename="{nome_file}.{formato}"'},
    )
