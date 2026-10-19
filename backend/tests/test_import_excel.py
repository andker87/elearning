"""2.3.3: importazione da Excel o CSV con anteprima; uffici e mansioni mancanti non scartano la riga."""
import io

import pytest
from openpyxl import Workbook, load_workbook

from app.models.group import UserGroup
from app.models.user import User
from app.services.impronte import impronta


def invia(client, auth, admin, nome, contenuto, tipo="text/csv", **params):
    return client.post("/api/auth/users/import-csv", headers=auth(admin), params=params, files={"file": (nome, contenuto, tipo)})


def xlsx(righe, foglio="Utenti"):
    wb = Workbook()
    ws = wb.active
    ws.title = foglio
    for r in righe:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
INTEST = ["email", "nome", "cognome", "ufficio", "mansioni", "admin", "formazione", "formatore"]


# --- anteprima -------------------------------------------------------------------------------------

def test_l_anteprima_non_crea_niente_e_non_manda_email(client, auth, admin, db, crea_ufficio, monkeypatch):
    from test_inviti import canale
    inviate = canale(db, monkeypatch)
    crea_ufficio("Produzione")
    prima = db.query(User).count()
    testo = "email,nome,cognome,ufficio\nmario@az.it,Mario,Rossi,Produzione\nanna@az.it,Anna,Verdi,Inesistente\n"
    r = invia(client, auth, admin, "u.csv", testo, anteprima="true").json()
    assert r["anteprima"] is True and r["creati"] == 2 and r["con_avvisi"] == 1 and r["errori"] == 0
    assert "«Inesistente» non presente" in r["righe"][1]["avvisi"][0]
    assert all(x["password"] is None for x in r["righe"])                      # le password si mostrano solo a importazione fatta
    db.expire_all()
    assert db.query(User).count() == prima and inviate == []                  # niente persone, niente email
    vero = invia(client, auth, admin, "u.csv", testo).json()                    # la stessa chiamata senza anteprima importa davvero
    assert vero["anteprima"] is False and vero["creati"] == 2 and db.query(User).count() == prima + 2
    assert len(inviate) == 2


def test_anteprima_dei_cessati_non_disattiva(client, auth, admin, db, crea_utente):
    u = crea_utente("cessa@az.it")
    r = invia(client, auth, admin, "c.csv", "email\ncessa@az.it\nnonc'e@az.it\n", anteprima="true", modo="cessati").json()
    assert r["anteprima"] is True and r["disattivati"] == 1 and [x["esito"] for x in r["righe"]] == ["disattivato", "ignorato"]
    db.expire_all()
    assert db.get(User, u.id).is_active is True
    invia(client, auth, admin, "c.csv", "email\ncessa@az.it\n", modo="cessati")
    db.expire_all()
    assert db.get(User, u.id).is_active is False


def test_la_stessa_email_due_volte_nello_stesso_file(client, auth, admin, crea_ufficio):
    crea_ufficio("Produzione")
    r = invia(client, auth, admin, "u.csv", "email,nome,cognome,ufficio\na@az.it,A,A,Produzione\nA@az.it,B,B,Produzione\n").json()
    assert [x["esito"] for x in r["righe"]] == ["creato", "errore"] and "già esistente" in r["righe"][1]["messaggio"]


# --- Excel -----------------------------------------------------------------------------------------

def test_si_importa_un_file_excel_con_date_e_sino(client, auth, admin, db, crea_ufficio):
    from datetime import datetime
    crea_ufficio("Produzione")
    client.post("/api/mansioni", headers=auth(admin), json={"name": "Preposto"})
    contenuto = xlsx([INTEST + ["data_ingresso"],
                      ["mario@az.it", "Mario", "Rossi", "Produzione", "Preposto", "no", "sì", True, datetime(2026, 1, 15)],
                      [None, None, None, None, None, None, None, None, None],              # riga vuota: saltata
                      ["anna@az.it", "Anna", "Verdi", "Produzione", None, None, None, None, None]])
    r = invia(client, auth, admin, "utenti.xlsx", contenuto, XLSX).json()
    assert [x["esito"] for x in r["righe"]] == ["creato", "creato"] and [x["riga"] for x in r["righe"]] == [2, 4]   # i numeri sono quelli di Excel
    m = db.query(User).filter(User.email_idx == impronta("mario@az.it")).one()
    assert m.is_formazione and m.is_formatore and not m.is_admin and m.data_ingresso.date().isoformat() == "2026-01-15"


def test_l_excel_si_riconosce_dal_contenuto_anche_con_estensione_csv(client, auth, admin, crea_ufficio):
    crea_ufficio("Produzione")
    contenuto = xlsx([["email", "nome", "cognome", "ufficio"], ["x@az.it", "X", "Y", "Produzione"]])
    r = invia(client, auth, admin, "salvato-male.csv", contenuto)
    assert r.status_code == 200 and r.json()["creati"] == 1


def test_si_legge_il_foglio_utenti_e_non_il_primo(client, auth, admin, crea_ufficio):
    crea_ufficio("Produzione")
    wb = Workbook(); wb.active.title = "Istruzioni"; wb.active.append(["testo"])
    ws = wb.create_sheet("Utenti"); ws.append(["email", "nome", "cognome", "ufficio"]); ws.append(["x@az.it", "X", "Y", "Produzione"])
    buf = io.BytesIO(); wb.save(buf)
    assert invia(client, auth, admin, "u.xlsx", buf.getvalue(), XLSX).json()["creati"] == 1


@pytest.mark.parametrize("nome,contenuto,atteso", [
    ("vecchio.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1 resto", ".xls"),
    ("rotto.xlsx", b"PK\x03\x04 non e' uno zip vero", "Excel"),
    ("vuoto.csv", b"", "vuoto"),
])
def test_file_non_validi_spiegano_cosa_fare(client, auth, admin, nome, contenuto, atteso):
    r = invia(client, auth, admin, nome, contenuto)
    assert r.status_code == 400 and atteso in r.json()["detail"]


def test_un_csv_salvato_da_excel_in_codifica_windows_si_legge(client, auth, admin, db, crea_ufficio):
    crea_ufficio("Produzione")
    testo = "email;nome;cognome;ufficio\r\nx@az.it;Niccolò;D'Àngelo;Produzione\r\n".encode("cp1252")
    r = invia(client, auth, admin, "u.csv", testo).json()
    assert r["creati"] == 1
    u = db.query(User).filter(User.email_idx == impronta("x@az.it")).one()
    assert (u.nome, u.cognome) == ("Niccolò", "D'Àngelo")


def test_un_file_troppo_grande_e_rifiutato(client, auth, admin):
    r = invia(client, auth, admin, "u.csv", b"a" * (5 * 1024 * 1024 + 10))
    assert r.status_code == 400 and "troppo grande" in r.json()["detail"]


# --- modelli ---------------------------------------------------------------------------------------

def test_il_modello_excel_ha_i_nomi_veri_e_si_importa_cosi_com_e(client, auth, admin, db, crea_ufficio):
    padre = crea_ufficio("Produzione")
    from app.models.group import Group
    db.add(Group(name="Linea 2", parent_id=padre.id, is_ufficio=True)); db.commit()
    client.post("/api/mansioni", headers=auth(admin), json={"name": "Preposto"})
    r = client.get("/api/auth/users/import-modello?formato=xlsx&modo=nuovi", headers=auth(admin))
    assert r.status_code == 200 and "modello-nuovi-utenti.xlsx" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Utenti", "Istruzioni", "Uffici", "Mansioni"]
    assert [c.value for c in wb["Utenti"][1]][:4] == ["email", "nome", "cognome", "ufficio"]
    assert [c[0].value for c in wb["Uffici"].iter_rows(min_row=2)] == ["Produzione", "Produzione / Linea 2"]
    assert [c[0].value for c in wb["Mansioni"].iter_rows(min_row=2)] == ["Preposto"]
    esito = invia(client, auth, admin, "modello-nuovi-utenti.xlsx", r.content, XLSX).json()
    assert esito["creati"] == 1 and esito["righe"][0]["avvisi"] == []          # la riga di esempio e' valida


def test_il_modello_csv_e_il_modello_dei_cessati(client, auth, admin, crea_ufficio):
    crea_ufficio("Produzione")
    csv_ = client.get("/api/auth/users/import-modello?formato=csv&modo=nuovi", headers=auth(admin))
    testo = csv_.content.decode("utf-8-sig")
    assert testo.splitlines()[0] == "email;nome;cognome;ufficio;mansioni;admin;formazione;formatore;is_external;organizzazione;data_ingresso;password"
    assert "Produzione" in testo.splitlines()[1]
    assert invia(client, auth, admin, "m.csv", csv_.content).json()["creati"] == 1
    for formato in ("csv", "xlsx"):
        c = client.get(f"/api/auth/users/import-modello?formato={formato}&modo=cessati", headers=auth(admin))
        assert c.status_code == 200 and "modello-cessati" in c.headers["content-disposition"]
    assert client.get("/api/auth/users/import-modello?formato=pdf", headers=auth(admin)).status_code == 400


def test_il_modello_non_esegue_formule_nei_nomi(client, auth, admin, db, crea_ufficio):
    crea_ufficio("=HYPERLINK(\"http://x\")")
    wb = load_workbook(io.BytesIO(client.get("/api/auth/users/import-modello?formato=xlsx", headers=auth(admin)).content))
    cella = wb["Uffici"]["A2"]
    assert cella.data_type == "s" and cella.value.startswith("=")              # testo, non formula


def test_solo_l_admin_scarica_i_modelli_e_importa(client, auth, formazione, utente):
    for u in (formazione, utente):
        assert client.get("/api/auth/users/import-modello", headers=auth(u)).status_code == 403
        assert invia(client, auth, u, "u.csv", "email,nome,cognome\n").status_code == 403
